import docx

def create_dummy():
    doc = docx.Document()
    doc.add_heading('Ayushi - Software Engineer', 0)
    
    doc.add_heading('Summary', level=1)
    doc.add_paragraph('Results-oriented Software Engineer with 3 years of experience building Python and React applications. Strong skills in database design and API development.')
    
    doc.add_heading('Experience', level=1)
    p1 = doc.add_paragraph()
    p1.add_run('Software Engineer at TechCorp (2023 - Present)\n').bold = True
    p1.add_run('- Developed and optimized REST APIs using FastAPI, reducing response latency by 25%.\n')
    p1.add_run('- Led a team of 3 developers to migrate a legacy React dashboard to Next.js.\n')
    p1.add_run('- Implemented unit and integration tests using pytest, increasing coverage to 85%.')
    
    doc.add_heading('Skills', level=1)
    doc.add_paragraph('Technical: Python, React, JavaScript, SQL, FastAPI, Git\nSoft: Teamwork, Problem Solving, Communication')
    
    doc.save('dummy_resume.docx')
    print("Created dummy_resume.docx successfully!")

if __name__ == '__main__':
    create_dummy()
